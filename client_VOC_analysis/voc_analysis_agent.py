"""
11장 실습 - CrewAI 고객사 VOC 분석 Agent

실행:
    python voc_analysis_agent.py

환경 변수:
    OPENAI_API_KEY  필수(CrewAI 이슈 분석용, 미설정/오류 시 규칙 기반 분석으로 대체)
    VOC_LLM_MODEL   선택(기본값: openai/gpt-4o-mini)
"""

from __future__ import annotations

import json
import os
import re
import tempfile
import traceback
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Generator

import gradio as gr
import pandas as pd
import plotly.graph_objects as go
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image
from plotly.subplots import make_subplots
from wordcloud import WordCloud


APP_TITLE = "고객사 VOC 분석 Agent"
REQUIRED_COLUMNS = ["순번", "일자", "고객명", "산업군", "지역", "제품명", "분야", "불만"]
CATEGORY_COLUMNS = ["산업군", "제품명", "분야"]
KOREAN_FONT_NAME = "맑은 고딕"
REPORT_DIR = Path(tempfile.gettempdir()) / "voc_agent_reports"
REPORT_DIR.mkdir(parents=True, exist_ok=True)

STOPWORDS = {
    "그리고", "그러나", "하지만", "또한", "관련", "대한", "대해", "에서", "으로", "에게",
    "하는", "하고", "합니다", "입니다", "있습니다", "없습니다", "되다", "되어", "되는", "하다",
    "너무", "매우", "정말", "조금", "자주", "계속", "사용", "제품", "고객", "불만", "문제",
    "요청", "문의", "경우", "때문", "부분", "사항", "처리", "발생", "확인", "필요", "개선",
}

PALETTE = ["#73B9FF", "#FFD45C", "#FF9EB5", "#92D7B5", "#B9A7F8", "#FFB66E"]


def _log(lines: list[str], message: str) -> str:
    stamp = datetime.now().strftime("%H:%M:%S")
    lines.append(f"[{stamp}] {message}")
    return "\n".join(lines)


def _empty_figure(message: str = "CSV 파일을 먼저 업로드해 주세요.") -> go.Figure:
    fig = go.Figure()
    fig.add_annotation(text=message, x=0.5, y=0.5, xref="paper", yref="paper", showarrow=False,
                       font=dict(size=17, color="#6C6575"))
    fig.update_layout(template="plotly_white", height=430, margin=dict(l=25, r=25, t=45, b=25),
        font=dict(family="Jua, Gowun Dodum, Malgun Gothic, Apple SD Gothic Neo, Noto Sans KR, sans-serif"))
    return fig


def _state_to_df(state: dict[str, Any] | None) -> pd.DataFrame:
    if not state or not state.get("records"):
        raise ValueError("CSV 파일을 먼저 업로드해 주세요.")
    df = pd.DataFrame(state["records"])
    for col in REQUIRED_COLUMNS:
        if col not in df.columns:
            raise ValueError(f"필수 열이 없습니다: {col}")
    return df[REQUIRED_COLUMNS].copy()


def _read_csv_safely(file_path: str) -> tuple[pd.DataFrame, str]:
    errors: list[str] = []
    for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            df = pd.read_csv(file_path, encoding=encoding)
            df.columns = [str(c).strip() for c in df.columns]
            missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
            if missing:
                raise ValueError("필수 열 누락: " + ", ".join(missing))
            df = df[REQUIRED_COLUMNS].copy()
            for col in REQUIRED_COLUMNS:
                df[col] = df[col].fillna("").astype(str).str.strip()
            if df.empty:
                raise ValueError("CSV에 데이터 행이 없습니다.")
            parsed_dates = pd.to_datetime(df["일자"], errors="coerce")
            if parsed_dates.notna().any():
                df.loc[parsed_dates.notna(), "일자"] = parsed_dates[parsed_dates.notna()].dt.strftime("%Y-%m-%d")
            return df, encoding
        except UnicodeDecodeError as exc:
            errors.append(f"{encoding}: {exc}")
        except pd.errors.ParserError as exc:
            raise ValueError(f"CSV 형식 오류: {exc}") from exc
    raise ValueError("CSV 인코딩을 읽지 못했습니다. UTF-8 또는 CP949로 저장해 주세요. " + " | ".join(errors))


def upload_csv(file_path: str | None) -> Generator[tuple[dict[str, Any] | None, pd.DataFrame, str, str], None, None]:
    logs: list[str] = []
    if not file_path:
        yield None, pd.DataFrame(columns=REQUIRED_COLUMNS), "파일을 선택해 주세요.", _log(logs, "업로드 대기")
        return
    yield None, pd.DataFrame(columns=REQUIRED_COLUMNS), "CSV 파일을 읽는 중입니다…", _log(logs, "파일 로드 시작")
    try:
        df, encoding = _read_csv_safely(file_path)
        state = {"records": df.to_dict(orient="records"), "source_name": Path(file_path).name}
        _log(logs, f"인코딩 확인: {encoding}")
        _log(logs, f"스키마 검증 완료: {len(df):,}건 × {len(df.columns)}개 열")
        yield state, df, f"✅ {len(df):,}건의 VOC를 불러왔습니다.", "\n".join(logs)
    except Exception as exc:
        _log(logs, f"오류: {exc}")
        yield None, pd.DataFrame(columns=REQUIRED_COLUMNS), f"❌ {exc}", "\n".join(logs)


def _ratio_table(df: pd.DataFrame, column: str) -> pd.DataFrame:
    values = df[column].replace("", "미입력").fillna("미입력")
    counts = values.value_counts(dropna=False)
    result = counts.rename_axis(column).reset_index(name="건수")
    result["비율(%)"] = (result["건수"] / max(len(df), 1) * 100).round(2)
    return result


def build_statistics_figure(df: pd.DataFrame, selected: list[str]) -> tuple[go.Figure, pd.DataFrame]:
    selected = [c for c in selected if c in CATEGORY_COLUMNS]
    if not selected:
        raise ValueError("산업군, 제품명, 분야 중 하나 이상을 선택해 주세요.")
    fig = make_subplots(rows=len(selected), cols=1, subplot_titles=[f"{c}별 VOC 비율" for c in selected],
                        vertical_spacing=max(0.10, 0.18 / len(selected)))
    combined: list[pd.DataFrame] = []
    for idx, col in enumerate(selected, start=1):
        table = _ratio_table(df, col)
        display = table.sort_values("비율(%)", ascending=True)
        fig.add_trace(
            go.Bar(
                x=display["비율(%)"], y=display[col], orientation="h",
                text=[f"{v:.2f}%" for v in display["비율(%)"]], textposition="outside",
                marker_color=PALETTE[(idx - 1) % len(PALETTE)], hovertemplate="%{y}<br>%{x:.2f}%<extra></extra>",
                name=col, showlegend=False,
            ), row=idx, col=1,
        )
        fig.update_xaxes(title_text="비율(%)", ticksuffix="%", row=idx, col=1, rangemode="tozero")
        table.insert(0, "구분", col)
        table = table.rename(columns={col: "항목"})
        combined.append(table[["구분", "항목", "건수", "비율(%)"]])
    fig.update_layout(
        template="plotly_white", height=max(430, 330 * len(selected)), bargap=0.28,
        margin=dict(l=65, r=75, t=70, b=45), paper_bgcolor="#FFFDF8", plot_bgcolor="#FFFFFF",
        font=dict(family="Jua, Gowun Dodum, Malgun Gothic, Apple SD Gothic Neo, Noto Sans KR, sans-serif", size=13, color="#3D3645"),
        title=dict(text=f"총 {len(df):,}건 기준 구성비", x=0.02, font=dict(size=20)),
    )
    return fig, pd.concat(combined, ignore_index=True)


def analyze_statistics(state: dict[str, Any] | None, selected: list[str]) -> Generator[tuple[go.Figure, pd.DataFrame, str], None, None]:
    logs: list[str] = []
    yield _empty_figure("통계 계산 중…"), pd.DataFrame(), _log(logs, "VOC 통계 분석 시작")
    try:
        df = _state_to_df(state)
        fig, table = build_statistics_figure(df, selected or [])
        _log(logs, f"선택 항목 분석 완료: {', '.join(selected)}")
        _log(logs, "Plotly 비율 막대그래프 생성 완료(소수점 둘째 자리)")
        yield fig, table, "\n".join(logs)
    except Exception as exc:
        yield _empty_figure(str(exc)), pd.DataFrame(), _log(logs, f"오류: {exc}")


def _find_korean_font() -> str:
    candidates = [
        os.environ.get("VOC_KOREAN_FONT", ""),
        r"C:\Windows\Fonts\malgun.ttf",
        r"C:\Windows\Fonts\malgunsl.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return candidate
    raise FileNotFoundError("한글 폰트를 찾지 못했습니다. VOC_KOREAN_FONT 환경변수에 폰트 경로를 지정해 주세요.")


def extract_keywords(texts: pd.Series, top_n: int = 120) -> Counter[str]:
    words: list[str] = []
    for text in texts.fillna("").astype(str):
        tokens = re.findall(r"[가-힣]{2,}|[A-Za-z][A-Za-z0-9_+.-]{1,}", text.lower())
        normalized: list[str] = []
        for token in tokens:
            # 형태소 분석기 없이도 워드클라우드의 '오류가/오류를' 같은 분산을 줄이는 보수적 조사 제거
            if re.fullmatch(r"[가-힣]+", token):
                for suffix in ("에게서", "으로", "에서", "에게", "부터", "까지", "은", "는", "이", "가", "을", "를", "와", "과", "의", "에", "로", "도", "만"):
                    if token.endswith(suffix) and len(token) - len(suffix) >= 2:
                        token = token[:-len(suffix)]
                        break
            if token not in STOPWORDS and not token.isdigit() and len(token) >= 2:
                normalized.append(token)
        words.extend(normalized)
    frequencies = Counter(words)
    return Counter(dict(frequencies.most_common(top_n)))


def make_wordcloud_file(df: pd.DataFrame) -> tuple[str, Counter[str]]:
    frequencies = extract_keywords(df["불만"])
    if not frequencies:
        raise ValueError("불만 열에서 워드클라우드용 키워드를 찾지 못했습니다.")
    font_path = _find_korean_font()
    output_path = REPORT_DIR / f"wordcloud_{datetime.now():%Y%m%d_%H%M%S_%f}.png"
    cloud = WordCloud(
        font_path=font_path, width=1400, height=760, background_color="#FFFDF8",
        colormap="tab20", max_words=120, prefer_horizontal=0.88, relative_scaling=0.45,
        collocations=False, random_state=42, margin=6,
    ).generate_from_frequencies(frequencies)
    cloud.to_file(str(output_path))
    return str(output_path), frequencies


def generate_wordcloud(state: dict[str, Any] | None) -> Generator[tuple[str | None, pd.DataFrame, str], None, None]:
    logs: list[str] = []
    yield None, pd.DataFrame(), _log(logs, "불만 텍스트 정제 및 키워드 추출 시작")
    try:
        df = _state_to_df(state)
        image_path, frequencies = make_wordcloud_file(df)
        top = pd.DataFrame(frequencies.most_common(30), columns=["키워드", "빈도"])
        _log(logs, f"키워드 {len(frequencies):,}개 추출")
        _log(logs, f"한글 폰트 적용 및 워드클라우드 생성 완료: {_find_korean_font()}")
        yield image_path, top, "\n".join(logs)
    except Exception as exc:
        yield None, pd.DataFrame(), _log(logs, f"오류: {exc}")


def _deterministic_issues(df: pd.DataFrame) -> dict[str, Any]:
    result: list[dict[str, Any]] = []
    industries = df["산업군"].replace("", "미입력").value_counts()
    for industry, count in industries.items():
        subset = df[df["산업군"].replace("", "미입력") == industry]
        field = subset["분야"].replace("", "미입력").value_counts().index[0]
        keywords = extract_keywords(subset["불만"], top_n=5)
        keyword_text = ", ".join(keywords.keys()) or "세부 키워드 없음"
        issue = f"{field} 관련 VOC 집중(대표 키워드: {keyword_text})"
        action = f"{field} 담당 조직이 대표 VOC를 원인 유형별로 재분류하고, 상위 원인의 조치 기준·담당자·완료 목표일을 확정"
        result.append({
            "산업군": str(industry), "건수": int(count), "주요이슈": issue,
            "개선과제": action, "근거": f"해당 산업군 {count}건 중 최다 분야는 {field}",
        })
    return {
        "산업군별_이슈": result,
        "공통_대응방안": [
            "상위 VOC 유형의 접수-분류-담당-완료 상태를 주 단위로 관리",
            "반복 VOC는 제품·운영·고객지원으로 원인을 구분하고 재발률을 추적",
            "개선 전후 VOC 건수와 처리기간을 동일 기준으로 비교",
        ],
        "analysis_mode": "규칙 기반",
    }


def _compact_dataset_for_llm(df: pd.DataFrame) -> str:
    industry = _ratio_table(df, "산업군").to_dict(orient="records")
    field = _ratio_table(df, "분야").to_dict(orient="records")
    cross = pd.crosstab(df["산업군"].replace("", "미입력"), df["분야"].replace("", "미입력"))
    samples = df[["산업군", "제품명", "분야", "불만"]].head(80).to_dict(orient="records")
    payload = {"총건수": len(df), "산업군통계": industry, "분야통계": field,
               "산업군_분야_교차표": cross.to_dict(), "VOC표본": samples}
    return json.dumps(payload, ensure_ascii=False)


def _extract_json(text: str) -> dict[str, Any]:
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip(), flags=re.IGNORECASE | re.DOTALL)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, flags=re.DOTALL)
        if not match:
            raise
        return json.loads(match.group(0))


def run_crewai_analysis(df: pd.DataFrame) -> tuple[dict[str, Any], str, str]:
    """voc → issue → report 세 Agent를 순차 실행한다. 실패 시 호출자가 대체 분석을 사용한다."""
    if not os.getenv("OPENAI_API_KEY"):
        raise RuntimeError("OPENAI_API_KEY가 없어 규칙 기반 분석을 사용합니다.")
    from crewai import Agent, Crew, LLM, Process, Task

    llm = LLM(
        model=os.getenv("VOC_LLM_MODEL", "openai/gpt-4o-mini"),
        api_key=os.getenv("OPENAI_API_KEY"), temperature=0.2,
    )
    voc_agent = Agent(
        role="VOC 데이터 분석가", goal="제공된 VOC 통계에서 사실과 집중 구간을 정확히 요약한다.",
        backstory="수치를 임의로 만들지 않고 입력 데이터의 비교·분포를 분석하는 시니어 분석가다.",
        llm=llm, verbose=True, allow_delegation=False,
    )
    issue_agent = Agent(
        role="VOC 이슈 전략가", goal="산업군별 주요 이슈와 실행 가능한 개선 과제를 근거와 함께 도출한다.",
        backstory="Fact-이슈-권고를 연결하고, 상관관계를 인과관계로 과장하지 않는 고객경험 전문가다.",
        llm=llm, verbose=True, allow_delegation=False,
    )
    report_agent = Agent(
        role="경영 보고서 작성자", goal="의사결정자가 5분 안에 판단할 수 있는 한국어 요약을 작성한다.",
        backstory="결론부터 쓰고 사실·해석·권고를 구분하는 한국어 업무보고 전문가다.",
        llm=llm, verbose=True, allow_delegation=False,
    )

    data_text = _compact_dataset_for_llm(df)
    voc_task = Task(
        description=("아래 VOC 데이터 요약을 분석하라. 산업군/분야 집중도, 대표 키워드, 주의할 데이터 한계를 "
                     "한국어로 정리하되 입력에 없는 수치를 만들지 마라.\n\n" + data_text),
        expected_output="근거 수치를 포함한 한국어 VOC 분석 메모", agent=voc_agent,
    )
    issue_task = Task(
        description=("VOC 분석을 바탕으로 산업군별 주요 이슈와 개선 과제를 도출하라. 반드시 다음 JSON만 출력하라: "
                     '{"산업군별_이슈":[{"산업군":"", "건수":0, "주요이슈":"", "개선과제":"", "근거":""}], '
                     '"공통_대응방안":["..."]}. 건수는 입력 데이터에 있는 값만 사용하라.'),
        expected_output="유효한 JSON 객체", agent=issue_agent, context=[voc_task],
    )
    report_task = Task(
        description=("앞선 분석과 이슈를 사용해 경영진용 요약을 작성하라. 구성은 ①한 줄 결론 ②핵심 사실 3개 "
                     "③우선 실행 과제 3개 ④데이터 한계다. 과장하거나 근거 없는 원인을 단정하지 마라."),
        expected_output="600자 이내 한국어 경영진 요약", agent=report_agent, context=[voc_task, issue_task],
    )
    crew = Crew(agents=[voc_agent, issue_agent, report_agent], tasks=[voc_task, issue_task, report_task],
                process=Process.sequential, verbose=True)
    crew.kickoff()
    issue_data = _extract_json(issue_task.output.raw)
    issue_data["analysis_mode"] = "CrewAI + OpenAI"
    return issue_data, report_task.output.raw.strip(), voc_task.output.raw.strip()


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _set_run_font(run, name: str = KOREAN_FONT_NAME, size: float | None = None,
                  bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        _set_run_font(run, size=16 if level == 1 else 13, bold=True, color="3D3645")


def _add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    _set_run_font(run, size=10)


def _add_dataframe_table(doc: Document, table_df: pd.DataFrame, max_rows: int = 20) -> None:
    view = table_df.head(max_rows).copy()
    table = doc.add_table(rows=1, cols=len(view.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    _set_repeat_table_header(header)
    for idx, col in enumerate(view.columns):
        cell = header.cells[idx]
        cell.text = str(col)
        _set_cell_shading(cell, "73B9FF")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            _set_run_font(run, size=9, bold=True, color="FFFFFF")
    for _, row in view.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[idx].paragraphs[0].runs:
                _set_run_font(run, size=9)


def _write_plotly_png(fig: go.Figure, name: str) -> str | None:
    path = REPORT_DIR / f"{name}_{datetime.now():%Y%m%d_%H%M%S_%f}.png"
    try:
        fig.write_image(str(path), width=1200, height=max(600, int(fig.layout.height or 600)), scale=1.5)
        return str(path)
    except Exception:
        return None


def create_word_report(df: pd.DataFrame, issue_data: dict[str, Any], executive_summary: str,
                       voc_summary: str, wordcloud_path: str) -> str:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = KOREAN_FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), KOREAN_FONT_NAME)
    normal.font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(20)
    title.paragraph_format.space_after = Pt(8)
    _set_run_font(title.add_run("고객사 VOC 분석 보고서"), size=24, bold=True, color="3D3645")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(subtitle.add_run(f"분석 대상 {len(df):,}건  |  생성일 {datetime.now():%Y-%m-%d}"), size=11, color="716879")

    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.cell(0, 0)
    _set_cell_shading(cell, "FFF3C4")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(7)
    _set_run_font(p.add_run("의사결정 요약\n"), size=11, bold=True, color="3D3645")
    _set_run_font(p.add_run(executive_summary or "상위 VOC 집중 구간을 우선 점검하고 개선 전후 지표를 동일 기준으로 추적합니다."), size=10)

    _add_heading(doc, "1. 분석 범위와 기준", 1)
    for text in [
        f"분석 대상: 업로드 CSV의 VOC {len(df):,}건",
        "분석 차원: 산업군, 제품명, 분야의 건수 및 구성비",
        f"이슈 도출 방식: {issue_data.get('analysis_mode', '분석')}",
        "주의: VOC 빈도는 고객 영향도나 원인 관계를 직접 증명하지 않으며, 우선순위 판단 시 심각도·재발률·처리기간을 함께 확인해야 합니다.",
    ]:
        _add_bullet(doc, text)

    industry_table = _ratio_table(df, "산업군")
    field_table = _ratio_table(df, "분야")
    product_table = _ratio_table(df, "제품명")
    _add_heading(doc, "2. 산업군별 VOC 통계", 1)
    _add_dataframe_table(doc, industry_table)
    fig_industry, _ = build_statistics_figure(df, ["산업군"])
    img = _write_plotly_png(fig_industry, "industry")
    if img:
        doc.add_picture(img, width=Inches(6.3))

    _add_heading(doc, "3. 분야별 VOC 통계", 1)
    _add_dataframe_table(doc, field_table)
    fig_field, _ = build_statistics_figure(df, ["분야"])
    img = _write_plotly_png(fig_field, "field")
    if img:
        doc.add_picture(img, width=Inches(6.3))

    _add_heading(doc, "4. 제품별 VOC 통계", 1)
    _add_dataframe_table(doc, product_table)

    _add_heading(doc, "5. 불만 키워드 워드클라우드", 1)
    doc.add_picture(wordcloud_path, width=Inches(6.4))
    caption = doc.add_paragraph("※ 불만 텍스트의 2자 이상 한글·영문 토큰 빈도 기준이며, 조사·일반어를 일부 제외했습니다.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        _set_run_font(run, size=8, color="716879")

    doc.add_section(WD_SECTION.NEW_PAGE)
    _add_heading(doc, "6. 산업군별 주요 이슈와 개선 과제", 1)
    issue_rows = issue_data.get("산업군별_이슈", [])
    if issue_rows:
        issue_df = pd.DataFrame(issue_rows)
        desired = [c for c in ["산업군", "건수", "주요이슈", "개선과제", "근거"] if c in issue_df.columns]
        _add_dataframe_table(doc, issue_df[desired], max_rows=30)
    else:
        _add_bullet(doc, "도출된 산업군별 이슈가 없습니다.")

    _add_heading(doc, "7. 공통 대응 방안", 1)
    for action in issue_data.get("공통_대응방안", []):
        _add_bullet(doc, str(action))

    if voc_summary:
        _add_heading(doc, "부록. VOC 분석 메모", 1)
        for paragraph_text in [x.strip() for x in voc_summary.splitlines() if x.strip()]:
            p = doc.add_paragraph()
            _set_run_font(p.add_run(paragraph_text), size=9)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(footer.add_run("CrewAI VOC Analysis Agent · 내부 분석용"), size=8, color="938B99")

    output = REPORT_DIR / f"VOC_분석보고서_{datetime.now():%Y%m%d_%H%M%S}.docx"
    doc.save(str(output))
    return str(output)


def generate_report(state: dict[str, Any] | None) -> Generator[tuple[str | None, str, str, str], None, None]:
    logs: list[str] = []
    yield None, "", _log(logs, "report agent 시작"), "보고서를 준비하고 있습니다…"
    try:
        df = _state_to_df(state)
        _log(logs, f"분석 데이터 확인: {len(df):,}건")
        yield None, "", "\n".join(logs), "통계와 워드클라우드를 생성하고 있습니다…"
        wordcloud_path, _ = make_wordcloud_file(df)
        _log(logs, "워드클라우드 생성 완료")

        yield None, "", "\n".join(logs), "CrewAI가 산업군별 이슈와 개선 과제를 분석하고 있습니다…"
        try:
            issue_data, executive_summary, voc_summary = run_crewai_analysis(df)
            _log(logs, "voc → issue → report Agent 순차 협업 완료")
        except Exception as crew_exc:
            issue_data = _deterministic_issues(df)
            top_industry = _ratio_table(df, "산업군").iloc[0]
            top_field = _ratio_table(df, "분야").iloc[0]
            executive_summary = (
                f"총 {len(df):,}건 중 산업군은 '{top_industry['산업군']}'({top_industry['비율(%)']:.2f}%), "
                f"분야는 '{top_field['분야']}'({top_field['비율(%)']:.2f}%)의 비중이 가장 높습니다. "
                "상위 집중 구간의 원인을 재분류하고 담당자·완료 목표일·재발률을 함께 관리해야 합니다."
            )
            voc_summary = ""
            _log(logs, f"CrewAI 호출 대체 처리: {crew_exc}")
            _log(logs, "규칙 기반 이슈·개선 과제로 보고서를 계속 생성합니다.")

        yield None, executive_summary, "\n".join(logs), "Word 보고서에 표·그래프·이슈를 편집하고 있습니다…"
        report_path = create_word_report(df, issue_data, executive_summary, voc_summary, wordcloud_path)
        _log(logs, f"한글 Word 보고서 저장 완료: {Path(report_path).name}")
        preview = executive_summary + "\n\n" + "\n".join(
            f"• {x.get('산업군', '')}: {x.get('주요이슈', '')}\n  ↳ {x.get('개선과제', '')}"
            for x in issue_data.get("산업군별_이슈", [])
        )
        yield report_path, preview, "\n".join(logs), "✅ Word 보고서 생성이 완료되었습니다."
    except Exception as exc:
        _log(logs, f"보고서 생성 실패: {exc}")
        _log(logs, traceback.format_exc(limit=2))
        yield None, f"❌ {exc}", "\n".join(logs), "❌ 보고서 생성에 실패했습니다."


CUSTOM_CSS = """
@import url('https://fonts.googleapis.com/css2?family=Gowun+Dodum&family=Jua&display=swap');
:root { --cream:#fffaf0; --ink:#3d3645; --blue:#73b9ff; --yellow:#ffd45c; --pink:#ff9eb5; }
.gradio-container { background: linear-gradient(135deg,#fffaf0 0%,#fff7fb 48%,#f1f8ff 100%); color:var(--ink);
  font-family:'Gowun Dodum','Malgun Gothic','Apple SD Gothic Neo',sans-serif !important; }
.hero { position:relative; overflow:hidden; border:2px solid #eadfee; border-radius:26px; padding:22px 25px;
  background:rgba(255,255,255,.88); box-shadow:0 12px 35px rgba(91,75,105,.10); margin-bottom:14px; }
.hero h1 { margin:0; font-family:'Jua','Malgun Gothic',sans-serif; font-size:34px; font-weight:400;
  letter-spacing:.2px; color:var(--ink); text-shadow:2px 2px 0 #fff1a8; }
.hero p { margin:8px 0 0; color:#716879; font-size:15px; }
.hero img { position:absolute; right:0; top:0; width:290px; height:100%; object-fit:cover; object-position:center; opacity:.23; }
.badge { display:inline-block; background:#fff3c4; border:1px solid #f2d877; padding:5px 10px; border-radius:999px;
  font-family:'Jua','Malgun Gothic',sans-serif; font-size:13px; font-weight:400; margin-bottom:9px; }
.source-note {font-size:11px;color:#8d8493;margin:-7px 3px 12px;text-align:right}
.source-note a {color:#6e86ad}
.status-card { border-radius:15px !important; }
.logbox textarea { font-family:Consolas,'Malgun Gothic',monospace !important; font-size:12px !important; }
.primary { background:linear-gradient(135deg,#73b9ff,#9ba8ff) !important; border:none !important; }
button, [role='tab'], h1, h2, h3 { font-family:'Jua','Malgun Gothic',sans-serif !important; letter-spacing:.1px; }
button { font-size:16px !important; border-radius:16px !important; }
label, .prose, .markdown { font-family:'Gowun Dodum','Malgun Gothic',sans-serif !important; }
footer {display:none !important}
"""


def build_app() -> gr.Blocks:
    with gr.Blocks(title=APP_TITLE) as demo:
        state = gr.State(value=None)
        gr.HTML(
            """<div class="hero"><span class="badge">11장 실습 · CrewAI</span>
            <h1>고객사 VOC 분석 Agent</h1>
            <p>CSV 한 장으로 통계, 키워드, 산업군별 이슈와 실행 보고서까지.</p>
            <img src="https://www.anime-chiikawa.jp/images/ogp.jpg" alt="치이카와 공식 애니메이션 이미지"></div>
            <div class="source-note">캐릭터 이미지 © nagano / 교육용 UI · 출처: 
            <a href="https://www.anime-chiikawa.jp/" target="_blank">치이카와 공식 애니메이션 사이트</a></div>"""
        )

        with gr.Tabs():
            with gr.Tab("📁 파일 업로드"):
                gr.Markdown("### VOC CSV를 업로드하세요\n필수 항목: `순번, 일자, 고객명, 산업군, 지역, 제품명, 분야, 불만`")
                with gr.Row():
                    csv_file = gr.File(label="CSV 파일", file_types=[".csv"], type="filepath", scale=2)
                    with gr.Column(scale=1):
                        upload_btn = gr.Button("파일 불러오기", variant="primary")
                        upload_status = gr.Markdown("업로드 대기 중")
                data_preview = gr.Dataframe(label="업로드 데이터", interactive=False, wrap=True, max_height=440)
                upload_log = gr.Textbox(label="진행 로그", lines=7, interactive=False, elem_classes="logbox")

            with gr.Tab("📊 통계 분석"):
                gr.Markdown("### 보고 싶은 항목을 선택하세요\n각 항목의 전체 VOC 대비 비율을 소수점 둘째 자리까지 표시합니다.")
                stat_columns = gr.CheckboxGroup(CATEGORY_COLUMNS, value=CATEGORY_COLUMNS, label="분석 항목")
                stat_btn = gr.Button("통계 분석 실행", variant="primary")
                stat_plot = gr.Plot(value=_empty_figure())
                stat_table = gr.Dataframe(label="통계표", interactive=False)
                stat_log = gr.Textbox(label="진행 로그", lines=6, interactive=False, elem_classes="logbox")

            with gr.Tab("☁️ 워드클라우드"):
                gr.Markdown("### 불만 키워드 워드클라우드\n한글 폰트를 자동 탐색하고 반복 빈도가 높은 핵심어를 시각화합니다.")
                wc_btn = gr.Button("워드클라우드 생성", variant="primary")
                with gr.Row():
                    wc_image = gr.Image(label="워드클라우드", type="filepath", height=470)
                    wc_table = gr.Dataframe(label="상위 키워드", interactive=False)
                wc_log = gr.Textbox(label="진행 로그", lines=6, interactive=False, elem_classes="logbox")

            with gr.Tab("📝 보고서 생성"):
                gr.Markdown("### CrewAI 분석 보고서\n`voc → issue → report` Agent가 산업군별 이슈와 개선 과제를 정리하고 Word 보고서를 만듭니다.")
                gr.Markdown("> 🔒 CrewAI 분석 시 고객명·지역은 제외하지만, VOC 표본의 산업군·제품명·분야·불만 내용은 OpenAI API로 전송됩니다. 사내 개인정보·기밀정보 정책을 확인하세요.")
                report_btn = gr.Button("Word 보고서 생성", variant="primary")
                report_status = gr.Markdown("보고서 생성 대기 중")
                report_preview = gr.Textbox(label="주요 이슈 미리보기", lines=13, interactive=False)
                report_file = gr.File(label="보고서 다운로드", interactive=False)
                report_log = gr.Textbox(label="진행 로그", lines=10, interactive=False, elem_classes="logbox")

        upload_btn.click(upload_csv, inputs=[csv_file], outputs=[state, data_preview, upload_status, upload_log])
        csv_file.upload(upload_csv, inputs=[csv_file], outputs=[state, data_preview, upload_status, upload_log])
        stat_btn.click(analyze_statistics, inputs=[state, stat_columns], outputs=[stat_plot, stat_table, stat_log])
        wc_btn.click(generate_wordcloud, inputs=[state], outputs=[wc_image, wc_table, wc_log])
        report_btn.click(generate_report, inputs=[state], outputs=[report_file, report_preview, report_log, report_status])
    return demo


if __name__ == "__main__":
    app = build_app()
    app.queue(default_concurrency_limit=2).launch(
        inbrowser=True,
        theme=gr.themes.Soft(),
        css=CUSTOM_CSS,
    )
